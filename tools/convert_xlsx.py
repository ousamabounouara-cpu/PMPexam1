"""Convert the '180 Question PMP' sheet of PMP_practice_Q_A.xlsx into questions.js.
Usage: python tools/convert_xlsx.py PMP_practice_Q_A.xlsx questions.js
"""
import json, re, sys
from openpyxl import load_workbook

src, dst = sys.argv[1], sys.argv[2]
ws = load_workbook(src, read_only=True)["180 Question PMP"]
rows = [r for r in ws.iter_rows(values_only=True)]

questions, case, case_title, cur, case_range = [], [], None, None, (0, 0)
for r in rows:
    a = (str(r[0]).strip() if r[0] is not None else "")
    if not a or a in ("Total Score",):
        continue
    if re.match(r"^Case Study", a):
        case_title, case, cur = a, [], None
        rng = re.search(r"Questions\s+(\d+)\s*[–-]\s*(\d+)", a)
        case_range = (int(rng.group(1)), int(rng.group(2))) if rng else (0, 0)
        continue
    m = re.match(r"^Question\s+(\d+)$", a)
    if m:
        cur = {"id": int(m.group(1)), "case": None, "caseTitle": None, "text": "",
               "options": [], "answer": str(r[2]).strip().upper()}
        if case_title and case_range[0] <= cur["id"] <= case_range[1]:
            cur["caseTitle"] = re.sub(r"\s*\(Questions.*\)$", "", case_title)
            cur["case"] = list(case)
        questions.append(cur)
        continue
    if cur is None:
        if case_title:
            case.append(a)
        continue
    om = re.match(r"^([A-F])\.\s*(.*)$", a, re.S)
    if om:
        cur["options"].append({"key": om.group(1), "text": om.group(2).strip()})
    elif not cur["options"]:
        cur["text"] = (cur["text"] + " " + a).strip()

# Question numbering in the sheet may restart per case study, so renumber globally.
for i, q in enumerate(questions, 1):
    q["id"] = i
    q["answer"] = re.sub(r"[^A-F]", "", q["answer"])
    q["multi"] = len(q["answer"]) > 1
with open(dst, "w", encoding="utf-8") as f:
    body = json.dumps(questions, ensure_ascii=False, indent=1)
    f.write("window.QUESTIONS = " + body + ";\n" if dst.endswith(".js") else body)
print(len(questions), "questions;", sum(q["multi"] for q in questions), "multi-select;",
      sum(1 for q in questions if len(q["options"]) < 4), "with <4 options")
